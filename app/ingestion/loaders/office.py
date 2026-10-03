import time
import os
import logfire


def _parse_docx_native(file_path: str) -> str:
    import docx
    doc = docx.Document(file_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # Also extract text from tables
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                table_texts.append(row_text)
    return "\n\n".join(paragraphs + table_texts)


def _parse_pptx_native(file_path: str) -> str:
    from pptx import Presentation
    prs = Presentation(file_path)
    slide_texts = []
    for i, slide in enumerate(prs.slides):
        slide_parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_parts.append(shape.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        slide_parts.append(row_text)
        if slide_parts:
            slide_texts.append(f"--- Slide {i + 1} ---\n" + "\n".join(slide_parts))
    return "\n\n".join(slide_texts)


def parse_office(file_path: str) -> str:
    """
    Parses Office documents (.docx, .pptx) using python-docx and python-pptx directly.
    Bypasses heavy OCR or unstructured partitions to eliminate segfaults and ensure speed.
    Logs duration to Logfire.
    """
    start_time = time.time()
    ext = os.path.splitext(file_path)[1].lower()

    with logfire.span("📄 Office Document Parsing", filename=file_path, extension=ext):
        try:
            if ext == ".docx":
                full_text = _parse_docx_native(file_path)
            elif ext == ".pptx":
                full_text = _parse_pptx_native(file_path)
            else:
                # Fallback to unstructured if unexpected extension
                from unstructured.partition.auto import partition
                elements = partition(filename=file_path)
                full_text = "\n".join([str(el) for el in elements if str(el).strip()])

            duration = time.time() - start_time
            if not full_text.strip():
                logfire.warning(f"⚠️ Empty text extracted from {file_path} in {duration:.3f}s")
            else:
                logfire.info(f"✅ Successfully parsed {len(full_text)} characters in {duration:.3f}s from {file_path}")

            return full_text
        except Exception as e:
            duration = time.time() - start_time
            logfire.error(f"❌ Office Parse Failed for {file_path} after {duration:.3f}s: {e}")
            raise e
